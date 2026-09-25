import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = Path(__file__).resolve().parent.parent

# Colores de la paleta HydroGuard
C_FOREST = RGBColor(6, 78, 59)      # #064e3b - Verde Bosque Profundo
C_EMERALD = RGBColor(16, 185, 129)  # #10b981 - Esmeralda Primario
C_EMERALD_BG = RGBColor(236, 253, 245) # #ecfdf5 - Fondo Suave Verde
C_DARK = RGBColor(15, 23, 42)       # #0f172a - Slate Dark
C_MUTED = RGBColor(100, 116, 139)   # #64748b - Gris Texto Secundario
C_WHITE = RGBColor(255, 255, 255)   # Blanco
C_CARD_BG = RGBColor(248, 250, 252) # #f8fafc - Fondo Tarjetas
C_ROSE = RGBColor(225, 29, 72)      # #e11d48 - Alerta Crítica
C_AMBER = RGBColor(217, 119, 6)     # #d97706 - Advertencia
C_VIOLET = RGBColor(139, 92, 246)   # #8b5cf6 - Premium

def add_header(slide, title_text, category_text="HYDROGUARD • HACKATHON PITCH"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_EMERALD

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = C_DARK

def add_card(slide, left, top, width, height, title, items, border_color=None, bg_color=C_CARD_BG, title_color=C_DARK, icon=""):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.color.rgb = RGBColor(226, 232, 240)
        shape.line.width = Pt(1)

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = f"{icon} {title}".strip()
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = title_color
    p0.space_after = Pt(8)

    for it in items:
        p = tf.add_paragraph()
        p.text = f"• {it}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_DARK
        p.space_after = Pt(4)

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_FOREST
    bg1.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(3.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "AGROTECH & IOT SOLUTIONS • PITCH DECK 2026"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_EMERALD
    p_badge.space_after = Pt(10)

    p_main = tf.add_paragraph()
    p_main.text = "HYDROGUARD"
    p_main.font.size = Pt(54)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(8)

    p_sub = tf.add_paragraph()
    p_sub.text = "Sistema de Monitoreo Inteligente con Asistencia Pasiva para Invernaderos Hidropónicos"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(209, 250, 229)
    p_sub.space_after = Pt(18)

    p_meta = tf.add_paragraph()
    p_meta.text = "Previniendo pérdidas productivas mediante alertas tempranas y recetas correctivas en tiempo real."
    p_meta.font.size = Pt(13)
    p_meta.font.italic = True
    p_meta.font.color.rgb = RGBColor(167, 243, 208)

    # =========================================================================
    # SLIDE 2: EL PROBLEMA (PARTE 1) - LA VULNERABILIDAD HIDROPÓNICA
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "El Problema: La Alta Fragilidad del Cultivo Hidropónico", "1. CONTEXTO & PROBLEMÁTICA")

    col_w = Inches(3.64)
    h_cards = Inches(4.8)
    top_pos = Inches(1.8)

    add_card(
        s2, Inches(0.8), top_pos, col_w, h_cards,
        "Sensibilidad Química Crítica",
        [
            "En sistemas hidropónicos NFT, las raíces no tienen suelo como amortiguador natural.",
            "Una caída de pH (ej. de 6.0 a 5.0) o salinidad excesiva bloquea la asimilación de nutrientes.",
            "Una cosecha completa puede arruinarse irreversiblemente en menos de 24 horas."
        ],
        border_color=C_ROSE,
        title_color=C_ROSE,
        icon="⚠️"
    )

    add_card(
        s2, Inches(4.84), top_pos, col_w, h_cards,
        "Detección Tardía & Manual",
        [
            "La mayoría de productores mide con tiras reactivas o lápices portátiles de forma esporádica.",
            "La marchitez en hojas o quemaduras radiculares se notan cuando el daño ya es severo.",
            "No existe trazabilidad de variaciones nocturnas o picos de temperatura en la solución."
        ],
        border_color=C_AMBER,
        title_color=C_AMBER,
        icon="⏳"
    )

    add_card(
        s2, Inches(8.88), top_pos, col_w, h_cards,
        "Dependencia Energética",
        [
            "Las bombas de recirculación de agua y oxigenación dependen 100% de la energía eléctrica.",
            "Un corte de luz sin aviso detiene el flujo y seca las raíces en 30 minutos.",
            "Falta de visibilidad sobre el nivel de combustible del generador o carga de baterías solares."
        ],
        border_color=RGBColor(59, 130, 246),
        title_color=RGBColor(29, 78, 216),
        icon="⚡"
    )

    # =========================================================================
    # SLIDE 3: EL PROBLEMA (PARTE 2) - POR QUÉ FALLAN LAS SOLUCIONES ACTUALES
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "¿Por qué las soluciones tradicionales no resuelven el problema?", "1. CONTEXTO & PROBLEMÁTICA")

    w_half = Inches(5.66)
    h_half = Inches(2.25)

    add_card(
        s3, Inches(0.8), Inches(1.8), w_half, h_half,
        "1. Automatizaciones 'Ciegas' Peligrosas",
        [
            "Sistemas que inyectan químicos automáticamente con electroválvulas suelen fallar por bombas trabadas.",
            "La sobredosificación automática de ácidos quema cultivos enteros sin supervisión humana."
        ],
        border_color=C_ROSE, title_color=C_ROSE, icon="❌"
    )

    add_card(
        s3, Inches(6.86), Inches(1.8), w_half, h_half,
        "2. Costos Prohibitivos de Hardware",
        [
            "Soluciones industriales importadas exigen inversiones superiores a los $10,000 USD.",
            "Inaccesible para más del 90% de los productores agrícolas de la región."
        ],
        border_color=C_ROSE, title_color=C_ROSE, icon="❌"
    )

    add_card(
        s3, Inches(0.8), Inches(4.35), w_half, h_half,
        "3. Alertas Frías que no Guían",
        [
            "Los sistemas convencionales solo avisan 'pH fuera de rango', sin indicar qué cantidad de buffer añadir.",
            "El productor pierde tiempo valioso calculando proporciones manualmente bajo presión."
        ],
        border_color=C_AMBER, title_color=C_AMBER, icon="⚠️"
    )

    add_card(
        s3, Inches(6.86), Inches(4.35), w_half, h_half,
        "4. Planes Rígidos 'Todo o Nada'",
        [
            "Obligan a comprar módulos para generadores o paneles solares que el productor básico no posee.",
            "Falta de una plataforma modular que escale junto con el crecimiento de la finca."
        ],
        border_color=C_AMBER, title_color=C_AMBER, icon="⚠️"
    )

    # =========================================================================
    # SLIDE 4: LA SOLUCIÓN HYDROGUARD
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Nuestra Solución: Monitoreo Integral + Asistencia Pasiva", "2. PROPUESTA DE VALOR")

    add_card(
        s4, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.8),
        "Monitoreo Continuo IoT",
        [
            "Sensores en Tanques: pH, EC Nutrientes, Nivel de Agua, Temp Líquido, Caudal, Oxígeno Disuelto.",
            "Microclima Invernadero: Temp Aire, Humedad, Radiación Solar, CO₂ y Flujo de Aire.",
            "Transmisión en tiempo real vía microcontroladores ESP32 de bajo costo."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="📡"
    )

    add_card(
        s4, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.8),
        "Filosofía de Asistencia Pasiva",
        [
            "El sistema NO acciona motores a ciegas: protege la inversión evitando accidentes de sobredosificación.",
            "Calcula la corrección matemática exacta y le receta al productor el paso a paso.",
            "Seguridad agronómica garantizada con control humano en el circuito."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="🛡️"
    )

    add_card(
        s4, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.8),
        "Modelo Modular por Suscripción",
        [
            "Plan Base: Monitoreo de 4 tanques + Microclima + Recordatorios de tareas.",
            "Plan Estándar: Suma monitoreo de Generador Diésel y combustible.",
            "Plan Premium: Suma banco de Baterías de Litio y Paneles Solares."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="👑"
    )

    # =========================================================================
    # SLIDE 5: EL MOTOR DE ASISTENCIA PASIVA (RECETAS EXACTAS)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "La Innovación: De una Alerta Fría a una Receta Agronómica", "2. TECNOLOGÍA CENTRAL")

    add_card(
        s5, Inches(0.8), Inches(1.8), Inches(5.66), Inches(4.8),
        "El Enfoque Tradicional (Inútil)",
        [
            "Alerta en pantalla: '¡PELIGRO! pH ácido (5.0 pH) en Tanque 1'.",
            "Consecuencia: El productor entra en pánico.",
            "Dilema: ¿Cuánto buffer pH+ agrego para 1000 Litros?",
            "Riesgo: Si agrega de más, sube el pH a 7.5 y bloquea el hierro, dañando la cosecha."
        ],
        border_color=C_ROSE, title_color=C_ROSE, icon="❌"
    )

    add_card(
        s5, Inches(6.86), Inches(1.8), Inches(5.66), Inches(4.8),
        "Con HydroGuard (Receta Asistida)",
        [
            "Diagnóstico: pH 5.0 en Tanque #1 (Capacidad: 1000 Litros).",
            "Dosis Calculada: 200.0 ml de Solución Buffer Incrementadora (KOH al 10%).",
            "Paso 1: Medir exactamente 200 ml en probeta plástica limpia.",
            "Paso 2: Diluir en un balde con 5L de agua previa.",
            "Paso 3: Verter en el retorno del tanque con recirculación activa.",
            "Paso 4: Esperar 30 min y el sistema verifica automáticamente la estabilización a 6.0 pH."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, bg_color=C_EMERALD_BG, icon="✅"
    )

    # =========================================================================
    # SLIDE 6: ARQUITECTURA TÉCNICA & STACK
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Arquitectura de Software & Hardware Robusta", "3. ARQUITECTURA TÉCNICA")

    w_arch = Inches(2.78)
    h_arch = Inches(4.8)

    add_card(
        s6, Inches(0.8), Inches(1.8), w_arch, h_arch,
        "1. IoT Hardware",
        [
            "Microcontrolador ESP32-HydroGuard.",
            "Sensores industriales de pH, EC, nivel ultrasónico y temperatura.",
            "Heartbeat seguro con API Key por hash SHA-256."
        ],
        icon="🔌"
    )

    add_card(
        s6, Inches(3.83), Inches(1.8), w_arch, h_arch,
        "2. Backend FastAPI",
        [
            "API RESTful en Python de alto rendimiento asíncrono.",
            "Ingesta de telemetría por lotes (/telemetry/ingest).",
            "Motor de Asistencia Pasiva (AlertEngine) con fórmulas dinámicas."
        ],
        icon="⚙️"
    )

    add_card(
        s6, Inches(6.86), Inches(1.8), w_arch, h_arch,
        "3. PostgreSQL DB",
        [
            "Modelos relacionales con SQLAlchemy.",
            "Particionamiento temporal de lecturas para millones de filas.",
            "Campo 'pasos_resolucion' para trazabilidad de recetas."
        ],
        icon="🗄️"
    )

    add_card(
        s6, Inches(9.89), Inches(1.8), w_arch, h_arch,
        "4. Frontend Vivo",
        [
            "Vanilla CSS/JS moderno de alto contraste para campo.",
            "Botón de planes cíclico reactivo.",
            "Panel de Acción con acordeón interactivo de recetas."
        ],
        icon="💻"
    )

    # =========================================================================
    # SLIDE 7: MODELO DE NEGOCIO Y SUSCRIPCIÓN
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Modelo de Negocio: SaaS Modular por Escalas", "4. NEGOCIO & ESCALABILIDAD")

    add_card(
        s7, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.8),
        "Plan Base ($19.99/mes)",
        [
            "Ideal para productores familiares.",
            "Monitoreo de hasta 4 tanques hidropónicos.",
            "Sensores de Microclima Invernadero.",
            "Alertas y Recetas de Asistencia Pasiva.",
            "Recordatorios de nutrición A+B."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="🌱"
    )

    add_card(
        s7, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.8),
        "Plan Estándar ($39.99/mes)",
        [
            "Para granjas con riesgo de corte eléctrico.",
            "Todo lo incluido en el Plan Base.",
            "Monitoreo de Generador Diésel.",
            "Nivel de combustible en tiempo real.",
            "Cálculo de autonomía restante en horas."
        ],
        border_color=C_AMBER, title_color=C_AMBER, icon="⚡"
    )

    add_card(
        s7, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.8),
        "Plan Premium ($59.99/mes)",
        [
            "Para instalaciones 100% autosuficientes.",
            "Todo lo del Plan Estándar.",
            "Monitoreo de Paneles Fotovoltaicos.",
            "Banco de Baterías de Litio (% carga).",
            "Rendimiento diario y eficiencia solar."
        ],
        border_color=C_VIOLET, title_color=C_VIOLET, icon="👑"
    )

    # =========================================================================
    # SLIDE 8: IMPACTO Y VALIDACIÓN EN CAMPO
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Impacto Económico y Retorno de Inversión (ROI)", "5. IMPACTO EN EL PRODUCTOR")

    add_card(
        s8, Inches(0.8), Inches(1.8), Inches(5.66), Inches(2.25),
        "Reducción del 85% en Pérdidas",
        [
            "Evita el desbalance de nutrientes y la pudrición radicular por calor o falta de oxígeno.",
            "Salvar una sola cosecha de lechuga/frutilla paga 3 años de suscripción."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="📉"
    )

    add_card(
        s8, Inches(6.86), Inches(1.8), Inches(5.66), Inches(2.25),
        "Ahorro de Insumos Químicos",
        [
            "Las recetas calculan la dosis exacta en mililitros según la capacidad real del tanque.",
            "Elimina el desperdicio por sobre-dosificación de buffers y fertilizantes."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="🧪"
    )

    add_card(
        s8, Inches(0.8), Inches(4.35), Inches(5.66), Inches(2.25),
        "Respuesta Inmediata en Segundos",
        [
            "El productor recibe la notificación y el paso a paso antes de que el cultivo sufra estrés.",
            "Monitoreo 24/7 sin necesidad de inspecciones manuales nocturnas."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="⏱️"
    )

    add_card(
        s8, Inches(6.86), Inches(4.35), Inches(5.66), Inches(2.25),
        "Adopción Tecnológica Sin Fricción",
        [
            "Sin manuales complejos ni calibraciones engorrosas.",
            "Interfaz pensada para el productor en el campo desde su teléfono móvil."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="📱"
    )

    # =========================================================================
    # SLIDE 9: PRÓXIMOS PASOS (ROADMAP)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Roadmap: De Prototipo a Escala Comercial", "6. VISIÓN DE FUTURO")

    w_rd = Inches(3.64)
    h_rd = Inches(4.8)

    add_card(
        s9, Inches(0.8), Inches(1.8), w_rd, h_rd,
        "Fase 1: MVP Funcional (HOY)",
        [
            "✅ Telemetría IoT en vivo con PostgreSQL.",
            "✅ Motor de Asistencia Pasiva con recetas dinámicas.",
            "✅ Dashboard interactivo con planes Base, Estándar y Premium.",
            "✅ Sincronización reactiva sin hardcodeos."
        ],
        border_color=C_EMERALD, title_color=C_FOREST, icon="🚀"
    )

    add_card(
        s9, Inches(4.84), Inches(1.8), w_rd, h_rd,
        "Fase 2: Notificaciones Directas",
        [
            "📲 Integración con bot de WhatsApp Business y Telegram.",
            "🔔 Envío de recetas de emergencia directamente al teléfono del operario de turno.",
            "📊 Reporte semanal automatizado de salud del cultivo."
        ],
        border_color=RGBColor(59, 130, 246), title_color=RGBColor(29, 78, 216), icon="📲"
    )

    add_card(
        s9, Inches(8.88), Inches(1.8), w_rd, h_rd,
        "Fase 3: IA Predictiva",
        [
            "🤖 Modelos predictivos de consumo de nutrientes según curvas de radiación solar.",
            "📈 Detección temprana de patógenos antes de que se manifiesten en el agua.",
            "🌐 Red de sensores federada para benchmarking agronómico."
        ],
        border_color=C_VIOLET, title_color=C_VIOLET, icon="🤖"
    )

    # =========================================================================
    # SLIDE 10: CIERRE / CALL TO ACTION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = C_FOREST
    bg10.line.fill.background()

    tbox10 = s10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(3.2))
    tf10 = tbox10.text_frame
    tf10.word_wrap = True

    p_c1 = tf10.paragraphs[0]
    p_c1.text = "HYDROGUARD"
    p_c1.font.size = Pt(48)
    p_c1.font.bold = True
    p_c1.font.color.rgb = C_WHITE
    p_c1.space_after = Pt(12)

    p_c2 = tf10.add_paragraph()
    p_c2.text = "Protegiendo el futuro de la agricultura hidropónica, una gota a la vez."
    p_c2.font.size = Pt(22)
    p_c2.font.color.rgb = RGBColor(209, 250, 229)
    p_c2.space_after = Pt(24)

    p_c3 = tf10.add_paragraph()
    p_c3.text = "¡Muchas Gracias! • Espacio para Preguntas y Demostración en Vivo"
    p_c3.font.size = Pt(16)
    p_c3.font.bold = True
    p_c3.font.color.rgb = C_EMERALD

    # Guardar presentación
    out_dir = Path(BASE_DIR)
    out_file = out_dir / "HydroGuard_Pitch_Hackathon.pptx"
    prs.save(str(out_file))
    print(f"[SUCCESS] Presentación generada exitosamente en: {out_file}")

if __name__ == "__main__":
    build_presentation()
